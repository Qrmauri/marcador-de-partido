from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title', 'Heading 1']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(15)

def cue(text):
    p = doc.add_paragraph()
    r = p.add_run('En pantalla: ' + text)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string('536273')

def say(text):
    doc.add_paragraph(text)

doc.add_paragraph('Guion para presentar Marcador en Vivo', 'Title')
say('Duración sugerida de 3 a 4 minutos. Lee los párrafos de narración y usa las indicaciones en cursiva para coordinar lo que muestras en pantalla. El video demuestra el inicio de sesión con JWT, la consulta del rol y la obtención de datos de la API.')
doc.add_paragraph('Preparación antes de grabar', 'Heading 1')
say('Con la API y el frontend ejecutándose, abre http://127.0.0.1:5500. Presiona F12, entra en Network o Red y selecciona Fetch/XHR antes de iniciar sesión. Ten preparado un usuario válido. Las instrucciones de esta sección no se leen en el video.')
doc.add_paragraph('1 Presentación', 'Heading 1')
cue('Mostrar el frontend con los partidos y la tabla de posiciones.')
say('Hola, en este video voy a mostrar cómo el frontend externo de mi proyecto Marcador en Vivo se conecta con la API de Django. Primero explicaré el inicio de sesión con JWT y cómo se obtiene el rol del usuario. Después mostraré una consulta de partidos y cómo la respuesta de la API se convierte en el marcador que vemos en pantalla.')
doc.add_paragraph('2 Inicio de sesión', 'Heading 1')
cue('Iniciar sesión. En Network, seleccionar token/ y mostrar el método, la URL y la respuesta JSON.')
say('Cuando ingreso mi usuario y contraseña, el frontend envía una petición POST a /api/token/. Los datos se envían en formato JSON. La API valida las credenciales y, si son correctas, devuelve dos tokens: access y refresh.')
say('El token access sirve para autenticar las peticiones. El token refresh permite obtener un nuevo access cuando el anterior vence.')
cue('Abrir Application o Aplicación, luego Session Storage y seleccionar el origen del frontend. Mostrar las claves access y refresh.')
say('El frontend guarda ambos tokens en sessionStorage del navegador. No guarda la contraseña.')
doc.add_paragraph('3 Consulta del rol', 'Heading 1')
cue('Volver a Network y seleccionar me/. Mostrar Authorization en Headers y el JSON de Response. Señalar el usuario y el rol en la interfaz.')
say('Después del inicio de sesión, el frontend realiza una petición GET a /api/me/. En el encabezado Authorization envía la palabra Bearer seguida del token access. La API verifica el token y responde con el identificador, el nombre de usuario y el rol.')
say('El frontend utiliza ese rol para mostrar los controles correspondientes. El planillero puede modificar marcadores y el administrador también puede crear y borrar partidos. Los permisos también se verifican en el backend.')

doc.add_page_break()
doc.add_paragraph('4 Obtención de partidos', 'Heading 1')
cue('Limpiar Network, pulsar el filtro En juego y seleccionar la petición partidos/. Mostrar la URL con los parámetros liga y estado.')
say('Ahora voy a consultar los partidos en juego. Al pulsar este filtro, el frontend envía una petición GET a /api/partidos/, con los parámetros liga y estado. El parámetro liga identifica la liga seleccionada y estado tiene el valor en_juego.')
cue('Mostrar Response y desplegar un partido dentro de results. Señalar local_nombre, visita_nombre, goles_local, goles_visita y estado.')
say('La API consulta la base de datos y devuelve una respuesta JSON paginada. Dentro de results están los partidos. Cada partido incluye los nombres de los equipos, los goles del local, los goles de la visita y su estado.')
cue('Mostrar la tarjeta del mismo partido en el frontend.')
say('El frontend toma esos campos y construye la tarjeta que vemos aquí. Los nombres y el resultado en pantalla corresponden a la respuesta de la API.')
doc.add_paragraph('5 Tabla de posiciones', 'Heading 1')
cue('Recargar la página y seleccionar la petición que termina en /tabla/. Mostrar su respuesta y luego la tabla en la interfaz.')
say('Para obtener la tabla de posiciones, el frontend consulta /api/ligas/, seguido del identificador de la liga y /tabla/. El backend calcula las posiciones usando los partidos finalizados y devuelve los equipos con sus estadísticas y puntos. El frontend presenta esa respuesta como una tabla.')
doc.add_paragraph('6 Cierre', 'Heading 1')
cue('Mostrar nuevamente la vista general del frontend.')
say('Con esto queda demostrado el recorrido completo: el usuario interactúa con el frontend, el frontend consulta la API y la respuesta se muestra en pantalla. También mostré cómo se obtiene el token JWT, dónde se guarda y cómo se consulta el rol del usuario.')
doc.add_paragraph('Explicación opcional del diagrama', 'Heading 1')
cue('Mostrar el PNG del diagrama de secuencia al inicio o antes del cierre.')
say('Este diagrama resume el recorrido. Las flechas representan las peticiones y las respuestas. El frontend se comunica con la API, y la API consulta la base de datos. Los tokens se guardan en el navegador y el access se utiliza para identificar al usuario en las peticiones.')
doc.core_properties.title = 'Guion para presentar Marcador en Vivo'
doc.core_properties.subject = 'Demostración del frontend externo conectado a la API Django'
out = Path(__file__).parent / 'Guion_video_Marcador_en_Vivo.docx'
doc.save(out)
print(out)
